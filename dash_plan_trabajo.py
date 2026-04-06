"""
DISPOWER — Tablero de Seguimiento | Dirección Comercial ZNI
Usa refresh_token (mismo patrón que el otro app de la empresa).

secrets.toml — sección [microsoft]:
  tenant_id, client_id, client_secret, refresh_token, file_path
"""
import streamlit as st
import pandas as pd
import plotly.graph_objects as go
import plotly.express as px
from datetime import datetime
import io, os, warnings
import requests
from openpyxl import load_workbook

warnings.filterwarnings("ignore")

st.set_page_config(page_title="DISPOWER · ZNI", page_icon="⚡", layout="wide", initial_sidebar_state="expanded")

EXCEL_LOCAL = "DISPOWER_Tareas_Streamlit.xlsx"
HEADER_ROW  = 3

AREA_COLOR = {"Gestión Social":"#1A5C3A","SAC":"#2E6DA4","BI & Analítica":"#E67E22","SUI & Subsidios":"#7D3C98","Dirección":"#3D3D3D"}
ESTADOS     = ["Sin iniciar","En curso","Cerrada","Bloqueada","Cancelada"]
PRIORIDADES = ["CRÍTICA","ALTA","MEDIA","BAJA"]
ESTADO_COL  = {"Sin iniciar":"#E74C3C","En curso":"#F39C12","Cerrada":"#27AE60","Bloqueada":"#D85A30","Cancelada":"#888888"}
ESTADO_BG   = {"Sin iniciar":"#FADBD8","En curso":"#FEF9E7","Cerrada":"#D6F0E3","Bloqueada":"#FDEBD0","Cancelada":"#F5F5F5"}
PRIO_COL    = {"CRÍTICA":"#C0392B","ALTA":"#E67E22","MEDIA":"#27AE60","BAJA":"#2E6DA4"}
PRIO_BG     = {"CRÍTICA":"#FADBD8","ALTA":"#FDEBD0","MEDIA":"#D6F0E3","BAJA":"#D6E8F7"}
COLS = ["ID","ÁREA","PROYECTO","TAREA","CATEGORÍA","RESPONSABLE","FECHA INICIO","FECHA LÍMITE","PRIORIDAD","AVANCE","ESTADO","OBSERVACIÓN","FECHA CIERRE","CERRADA POR"]

st.markdown("""<style>
.kpi{background:white;border-radius:10px;padding:16px 18px;border-left:5px solid #2E6DA4;box-shadow:0 1px 5px rgba(0,0,0,.08);margin-bottom:10px}
.kv{font-size:2rem;font-weight:700;margin:4px 0}.kl{font-size:.7rem;color:#888;text-transform:uppercase;letter-spacing:.05em}.ks{font-size:.7rem;color:#aaa;margin-top:2px}
.badge{display:inline-block;border-radius:20px;padding:3px 10px;font-size:.72rem;font-weight:600}
.binfo{background:#EBF5FB;border-left:4px solid #2E6DA4;border-radius:4px;padding:10px 14px;margin:8px 0;font-size:.86rem;color:#1B3A5C}
.bok{background:#D6F0E3;border-left:4px solid #27AE60;border-radius:4px;padding:10px 14px;margin:8px 0;font-size:.86rem;color:#1A5C3A}
.bwarn{background:#FEF9E7;border-left:4px solid #F39C12;border-radius:4px;padding:10px 14px;margin:8px 0;font-size:.86rem;color:#7D6608}
.berr{background:#FADBD8;border-left:4px solid #E74C3C;border-radius:4px;padding:10px 14px;margin:8px 0;font-size:.86rem;color:#922B21}
</style>""", unsafe_allow_html=True)

# ── AUTH ─────────────────────────────────────────────────────────────────────
def get_cfg():
    try:    return dict(st.secrets.get("microsoft", {}))
    except: return {}

@st.cache_data(ttl=3000, show_spinner=False)
def get_access_token():
    cfg = get_cfg()
    tid,cid,sec,rt = cfg.get("tenant_id",""),cfg.get("client_id",""),cfg.get("client_secret",""),cfg.get("refresh_token","")
    if not all([tid,cid,sec,rt]): return None, None
    try:
        r = requests.post(
            f"https://login.microsoftonline.com/{tid}/oauth2/v2.0/token",
            data={"grant_type":"refresh_token","client_id":cid,"client_secret":sec,
                  "refresh_token":rt,"scope":"https://graph.microsoft.com/Files.ReadWrite offline_access"},
            timeout=20)
        if r.status_code == 200:
            d = r.json(); return d.get("access_token"), d.get("refresh_token")
        return None, None
    except Exception: return None, None

def g_bytes(path, token):
    try:
        r = requests.get(f"https://graph.microsoft.com/v1.0/{path}",headers={"Authorization":f"Bearer {token}"},timeout=40)
        return r.content if r.status_code==200 and len(r.content)>100 else None
    except: return None

def g_put(path, token, data):
    try:
        r = requests.put(f"https://graph.microsoft.com/v1.0/{path}",
            headers={"Authorization":f"Bearer {token}","Content-Type":"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
            data=data,timeout=60)
        return r.status_code in [200,201]
    except: return False

def g_json(path, token):
    try:
        r = requests.get(f"https://graph.microsoft.com/v1.0/{path}",headers={"Authorization":f"Bearer {token}"},timeout=20)
        return r.json() if r.status_code==200 else None
    except: return None

def file_endpoint():
    fp = get_cfg().get("file_path", f"/{EXCEL_LOCAL}")
    if not fp.startswith("/"): fp = "/" + fp
    return f"me/drive/root:{fp}:/content"

@st.cache_data(ttl=25, show_spinner=False)
def download_excel(token):
    raw = g_bytes(file_endpoint(), token)
    return raw if (raw and raw[:4]==b"PK\x03\x04") else None

def upload_excel(data, token): return g_put(file_endpoint(), token, data)

# ── EXCEL ────────────────────────────────────────────────────────────────────
def parse_excel(raw):
    try:
        xl = pd.ExcelFile(io.BytesIO(raw))
        target = next((s for s in xl.sheet_names if any(k in s.lower() for k in ["tarea","cierre","task"])),xl.sheet_names[0])
        df = xl.parse(target,header=HEADER_ROW,dtype=str).fillna("")
        df.columns = [str(c).strip().replace("\n"," ") for c in df.columns]
        aliases = {"ÁREA":"ÁREA","AREA":"ÁREA","CATEGORÍA":"CATEGORÍA","CATEGORIA":"CATEGORÍA",
            "OBSERVACIÓN":"OBSERVACIÓN","OBSERVACION":"OBSERVACIÓN","OBSERVACIÓN / BLOQUEO":"OBSERVACIÓN",
            "FECHA LÍMITE":"FECHA LÍMITE","FECHA LIMITE":"FECHA LÍMITE","% AVANCE":"AVANCE","AVANCE %":"AVANCE"}
        df = df.rename(columns={c:aliases.get(c.upper().strip(),c) for c in df.columns})
        if "ID" not in df.columns: return None
        df = df[df["ID"].astype(str).str.match(r"T-\d+",na=False)].copy()
        for c in COLS:
            if c not in df.columns: df[c]=""
        df = df[COLS].reset_index(drop=True)
        return df if len(df)>0 else None
    except Exception: return None

def build_bytes(df, orig):
    """
    Actualiza SOLO los valores de las celdas de datos en el Excel original.
    NO toca: formato, colores, bordes, celdas de título, hojas de instrucciones.
    El truco: actualizar celda por celda solo donde hay datos, nunca borrar filas.
    """
    if orig:
        try:
            wb = load_workbook(io.BytesIO(orig))
            target = next(
                (s for s in wb.sheetnames if any(k in s.lower() for k in ["tarea","cierre","task"])),
                wb.sheetnames[0]
            )
            ws = wb[target]

            # 1. Localizar la fila de encabezado (contiene ID y TAREA)
            hrow = None
            for ri in range(1, min(10, ws.max_row + 1)):
                vals = [str(ws.cell(ri, c).value or "").upper().strip()
                        for c in range(1, ws.max_column + 1)]
                if "ID" in vals and "TAREA" in vals:
                    hrow = ri
                    break
            if hrow is None:
                raise ValueError("No se encontró la fila de encabezado")

            # 2. Mapear nombre de columna → índice de columna en Excel
            hdrs = [str(ws.cell(hrow, c).value or "").strip().replace("\n", " ")
                    for c in range(1, ws.max_column + 1)]
            rev = {
                "OBSERVACIÓN / BLOQUEO": "OBSERVACIÓN",
                "% AVANCE": "AVANCE",
                "FECHA LÍMITE": "FECHA LÍMITE",
            }
            col_idx = {}  # nombre_col_app → número_columna_excel (1-based)
            for ci, h in enumerate(hdrs, 1):
                canon = rev.get(h.upper().strip(), h.upper().strip())
                for c in COLS:
                    if c.upper() == canon:
                        col_idx[c] = ci
                        break

            # 3. Construir índice de filas existentes en el Excel: ID → número de fila
            id_col = col_idx.get("ID")
            existing_rows = {}  # "T-001" → fila_excel
            if id_col:
                for ri in range(hrow + 1, ws.max_row + 1):
                    cell_val = str(ws.cell(ri, id_col).value or "").strip()
                    if cell_val.startswith("T-"):
                        existing_rows[cell_val] = ri

            # 4. Actualizar filas existentes — SOLO el valor, sin tocar formato
            last_data_row = hrow
            for _, row in df.iterrows():
                task_id = str(row.get("ID", "")).strip()
                if task_id in existing_rows:
                    er = existing_rows[task_id]
                    for col_name, ci2 in col_idx.items():
                        val = str(row.get(col_name, "")).strip()
                        ws.cell(er, ci2).value = val if val else None
                    last_data_row = max(last_data_row, er)
                else:
                    # Tarea nueva: agregar al final, copiando estilo de la última fila de datos
                    new_row = last_data_row + 1
                    # Copiar estilo de la fila anterior (misma estructura)
                    from copy import copy
                    from openpyxl.styles import PatternFill
                    for cc in range(1, ws.max_column + 1):
                        src_cell = ws.cell(last_data_row, cc)
                        dst_cell = ws.cell(new_row, cc)
                        if src_cell.has_style:
                            dst_cell.font      = copy(src_cell.font)
                            dst_cell.fill      = copy(src_cell.fill)
                            dst_cell.border    = copy(src_cell.border)
                            dst_cell.alignment = copy(src_cell.alignment)
                    # Escribir valores de la nueva tarea
                    for col_name, ci2 in col_idx.items():
                        val = str(row.get(col_name, "")).strip()
                        ws.cell(new_row, ci2).value = val if val else None
                    existing_rows[task_id] = new_row
                    last_data_row = new_row

            # 5. Limpiar filas que ya no existen en el df (tareas eliminadas)
            ids_en_df = set(df["ID"].astype(str).str.strip().tolist())
            for task_id, er in existing_rows.items():
                if task_id not in ids_en_df:
                    for cc in range(1, ws.max_column + 1):
                        ws.cell(er, cc).value = None

            buf = io.BytesIO()
            wb.save(buf)
            return buf.getvalue()

        except Exception as e:
            st.warning(f"Usando Excel simple (error al preservar formato): {e}")

    # Fallback: Excel simple sin formato
    buf2 = io.BytesIO()
    df.to_excel(buf2, index=False, sheet_name="Tareas", engine="openpyxl")
    return buf2.getvalue()

# ── SESIÓN ───────────────────────────────────────────────────────────────────
for k,v in [("df",None),("raw",None),("token",None),("hist",[]),("loaded",False),("src","")]:
    if k not in st.session_state: st.session_state[k]=v

def badge(t,fg,bg): return f'<span class="badge" style="background:{bg};color:{fg}">{t}</span>'
def eb(e): return badge(e,ESTADO_COL.get(e,"#888"),ESTADO_BG.get(e,"#eee"))
def pb(p): return badge(p,PRIO_COL.get(p,"#888"),PRIO_BG.get(p,"#eee"))
def add_hist(m): st.session_state["hist"].insert(0,{"ts":datetime.now().strftime("%d/%m/%Y %H:%M"),"msg":m})
def next_id(df):
    nums=pd.to_numeric(df["ID"].str.extract(r"T-(\d+)")[0],errors="coerce").dropna()
    return f"T-{int(nums.max())+1:03d}" if len(nums) else "T-001"

def commit(df_new, msg):
    st.session_state["df"]=df_new.copy()
    tok=st.session_state.get("token")
    if tok:
        nb=build_bytes(df_new,st.session_state.get("raw"))
        with st.spinner("Guardando en OneDrive..."):
            ok=upload_excel(nb,tok)
        download_excel.clear()
        add_hist(("✅ " if ok else "⚠️ ")+msg)
        st.toast("Guardado en OneDrive ✅" if ok else "Error al subir",icon="✅" if ok else "⚠️")
    else:
        add_hist(f"💾 {msg}"); st.toast("Guardado en sesión",icon="💾")

if not st.session_state["loaded"]:
    tok_i,_ = get_access_token()
    if tok_i:
        st.session_state["token"]=tok_i
        raw_i=download_excel(tok_i)
        if raw_i:
            df_i=parse_excel(raw_i)
            if df_i is not None:
                st.session_state.update({"df":df_i,"raw":raw_i,"src":"onedrive"})
    if st.session_state["df"] is None and os.path.exists(EXCEL_LOCAL):
        with open(EXCEL_LOCAL,"rb") as f: raw_l=f.read()
        df_l=parse_excel(raw_l)
        if df_l is not None: st.session_state.update({"df":df_l,"raw":raw_l,"src":"local"})
    st.session_state["loaded"]=True

df=st.session_state["df"]; src=st.session_state["src"]

# ── SIDEBAR ──────────────────────────────────────────────────────────────────
with st.sidebar:
    st.markdown("### ⚡ DISPOWER"); st.markdown("**Dirección Comercial ZNI**"); st.markdown("---")
    pagina=st.radio("Navegación",["🏠 Dashboard","📋 Lista de Tareas","🗂️ Kanban","📊 Proyectos ZNI","➕ Nueva Tarea","📜 Historial","⚙️ Configuración"])
    st.markdown("---")
    em={"onedrive":"🟢","local":"🟡","":"🔴"}.get(src,"🔴")
    lb={"onedrive":"OneDrive conectado","local":"Excel local (repo)","":"Sin conexión"}.get(src,"—")
    sb={"onedrive":"Cambios en tiempo real","local":"Configura OneDrive en ⚙️","":"Ve a ⚙️ Configuración"}.get(src,"")
    st.markdown(f"**{em} {lb}**"); st.caption(sb)
    if df is not None: st.caption(f"{len(df)} tareas · {int((df['ESTADO']=='Cerrada').sum())} cerradas")
    c1,c2=st.columns(2)
    with c1:
        if st.button("🔄 Recargar",use_container_width=True):
            download_excel.clear(); get_access_token.clear(); st.session_state["loaded"]=False; st.rerun()
    with c2:
        if df is not None:
            bdl=io.BytesIO(); df.to_excel(bdl,index=False,sheet_name="Tareas",engine="openpyxl")
            st.download_button("⬇️ Excel",bdl.getvalue(),f"DISPOWER_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",use_container_width=True)

if df is None and pagina!="⚙️ Configuración":
    st.markdown('<div class="bwarn">No hay datos. Configura el token en ⚙️ o sube el Excel manualmente.</div>',unsafe_allow_html=True)
    upl=st.file_uploader("📂 Subir Excel",type=["xlsx","xls"])
    if upl:
        raw_u=upl.read(); df_u=parse_excel(raw_u)
        if df_u is not None:
            st.session_state.update({"df":df_u,"raw":raw_u,"src":"local"}); st.success(f"✅ {len(df_u)} tareas."); st.rerun()
        else: st.error("No se pudo leer el archivo.")
    st.stop()

# ── DASHBOARD ────────────────────────────────────────────────────────────────
if pagina=="🏠 Dashboard":
    st.markdown("## ⚡ Dashboard Ejecutivo")
    st.caption("Director: Javier Alejandro Agudelo · Plan Estratégico 6 Meses · ZNI / SISFV")
    total=len(df); cerradas=int((df["ESTADO"]=="Cerrada").sum()); criticas=int(((df["PRIORIDAD"]=="CRÍTICA")&(df["ESTADO"]!="Cerrada")).sum())
    en_curso=int((df["ESTADO"]=="En curso").sum()); bloq=int((df["ESTADO"]=="Bloqueada").sum())
    try:    avg_av=round(pd.to_numeric(df["AVANCE"],errors="coerce").fillna(0).mean())
    except: avg_av=0
    for col,(val,lbl,sub3,cc) in zip(st.columns(6),[
        (total,"Total tareas",f"{cerradas} cerradas","#2E6DA4"),(criticas,"Críticas pend.","requieren atención","#C0392B"),
        (en_curso,"En curso","activas ahora","#E67E22"),(bloq,"Bloqueadas","requieren decisión","#D85A30"),
        (f"{avg_av}%","Avance global","promedio plan","#27AE60"),(cerradas,"Completadas",f"{round(cerradas/total*100) if total else 0}% del plan","#27AE60")]):
        col.markdown(f'<div class="kpi" style="border-left-color:{cc}"><div class="kl">{lbl}</div><div class="kv" style="color:{cc}">{val}</div><div class="ks">{sub3}</div></div>',unsafe_allow_html=True)
    st.markdown("---")
    cl,cr=st.columns([2,1])
    with cl:
        st.markdown("**Avance por área**")
        rows_a=[]
        for area in AREA_COLOR:
            s=df[df["ÁREA"]==area]
            if not len(s): continue
            try:    av=round(pd.to_numeric(s["AVANCE"],errors="coerce").fillna(0).mean())
            except: av=0
            rows_a.append({"Área":area,"Avance":av})
        if rows_a:
            fig=go.Figure(go.Bar(x=[r["Área"] for r in rows_a],y=[r["Avance"] for r in rows_a],
                marker_color=[AREA_COLOR[r["Área"]] for r in rows_a],opacity=.85,text=[f"{r['Avance']}%" for r in rows_a],textposition="outside"))
            fig.add_hline(y=100,line_dash="dash",line_color="#27AE60",line_width=1.5,annotation_text="Meta 100%")
            fig.update_layout(height=260,margin=dict(l=0,r=0,t=10,b=0),plot_bgcolor="white",paper_bgcolor="white",yaxis=dict(title="% avance",range=[0,120]),showlegend=False,font=dict(family="Arial",size=11))
            st.plotly_chart(fig,use_container_width=True)
    with cr:
        st.markdown("**Distribución de estados**")
        ec=df["ESTADO"].value_counts()
        if len(ec):
            fig2=go.Figure(go.Pie(labels=ec.index.tolist(),values=ec.values.tolist(),marker_colors=[ESTADO_COL.get(e,"#888") for e in ec.index],hole=.5,textinfo="value+percent"))
            fig2.update_layout(height=260,margin=dict(l=0,r=0,t=10,b=10),paper_bgcolor="white",showlegend=True,legend=dict(orientation="h",y=-.15),font=dict(family="Arial",size=11))
            st.plotly_chart(fig2,use_container_width=True)
    st.markdown("---")
    area_cols=st.columns(len(AREA_COLOR))
    for col,(area,color) in zip(area_cols,AREA_COLOR.items()):
        s=df[df["ÁREA"]==area]; cerr=int((s["ESTADO"]=="Cerrada").sum())
        crit=int(((s["PRIORIDAD"]=="CRÍTICA")&(s["ESTADO"]!="Cerrada")).sum())
        try:    av2=round(pd.to_numeric(s["AVANCE"],errors="coerce").fillna(0).mean())
        except: av2=0
        alerta=f'<div style="font-size:.64rem;color:#C0392B">⚠ {crit} críticas</div>' if crit else ""
        col.markdown(f'<div style="background:white;border-radius:8px;padding:12px;border-top:3px solid {color};box-shadow:0 1px 4px rgba(0,0,0,.07);text-align:center">'
            f'<div style="font-size:.68rem;color:{color};font-weight:600;margin-bottom:4px">{area}</div>'
            f'<div style="font-size:1.6rem;font-weight:700;color:{color}">{av2}%</div>'
            f'<div style="font-size:.68rem;color:#aaa">{len(s)} tareas · {cerr} cerradas</div>{alerta}</div>',unsafe_allow_html=True)
    st.markdown("---"); st.markdown("**⚠️ Tareas críticas pendientes:**")
    crit_df=df[(df["PRIORIDAD"]=="CRÍTICA")&(df["ESTADO"].isin(["Sin iniciar","Bloqueada"]))]
    if len(crit_df):
        for _,r in crit_df.iterrows():
            c1,c2,c3,c4=st.columns([1,5,2,2])
            c1.markdown(f'<b style="color:#C0392B">{r["ID"]}</b>',unsafe_allow_html=True)
            c2.write(r["TAREA"][:80]+("…" if len(r["TAREA"])>80 else ""))
            c3.markdown(f'<small style="color:{AREA_COLOR.get(r["ÁREA"],"#888")}">{r["ÁREA"]}</small>',unsafe_allow_html=True)
            c4.markdown(eb(r["ESTADO"]),unsafe_allow_html=True)
    else: st.success("Sin tareas críticas pendientes.")

# ── LISTA DE TAREAS ──────────────────────────────────────────────────────────
elif pagina=="📋 Lista de Tareas":
    st.markdown("## 📋 Lista de Tareas")
    fa=st.selectbox("Área",["Todas"]+sorted(df["ÁREA"].unique().tolist()),key="la")
    c1f,c2f,c3f=st.columns(3)
    with c1f: fe=st.selectbox("Estado",["Todos"]+ESTADOS,key="le")
    with c2f: fp=st.selectbox("Prioridad",["Todas"]+PRIORIDADES,key="lp")
    with c3f: fb=st.text_input("Buscar",placeholder="Tarea, proyecto, responsable...",key="lb")
    filt=df.copy()
    if fa!="Todas": filt=filt[filt["ÁREA"]==fa]
    if fe!="Todos": filt=filt[filt["ESTADO"]==fe]
    if fp!="Todas": filt=filt[filt["PRIORIDAD"]==fp]
    if fb:
        mask=filt["TAREA"].str.contains(fb,case=False,na=False)|filt["PROYECTO"].str.contains(fb,case=False,na=False)|filt["RESPONSABLE"].str.contains(fb,case=False,na=False)
        filt=filt[mask]
    st.caption(f"{len(filt)} de {len(df)} tareas"); st.markdown("---")
    for pos,(idx,row) in enumerate(filt.iterrows()):
        try:    av_v=int(float(row.get("AVANCE","0") or 0))
        except: av_v=0
        area_c=AREA_COLOR.get(row["ÁREA"],"#888"); wk=f"t{pos}"
        with st.expander(f"**{row['ID']}** — {row['TAREA'][:85]}{'…' if len(row['TAREA'])>85 else ''}",key=f"exp_{wk}"):
            h1,h2,h3,h4=st.columns(4)
            h1.markdown(eb(row["ESTADO"]),unsafe_allow_html=True); h2.markdown(pb(row["PRIORIDAD"]),unsafe_allow_html=True)
            h3.markdown(f'<span style="color:{area_c};font-weight:600;font-size:.8rem">{row["ÁREA"]}</span>',unsafe_allow_html=True)
            h4.markdown(f'<span style="color:#888;font-size:.8rem">{row["RESPONSABLE"]}</span>',unsafe_allow_html=True)
            prog_c="#27AE60" if av_v>=80 else ("#F39C12" if av_v>=40 else "#E74C3C")
            st.markdown(f'<div style="height:4px;background:#eee;border-radius:2px;margin:6px 0"><div style="height:4px;background:{prog_c};width:{av_v}%;border-radius:2px"></div></div>'
                f'<div style="font-size:.72rem;color:#aaa">Avance: {av_v}% · {row["PROYECTO"]} · {row["CATEGORÍA"]}</div>',unsafe_allow_html=True)
            if row.get("OBSERVACIÓN","").strip(): st.markdown(f'<div class="bwarn">{row["OBSERVACIÓN"]}</div>',unsafe_allow_html=True)
            st.markdown("---")
            e1,e2,e3=st.columns(3)
            with e1: new_est=st.selectbox("Estado",ESTADOS,index=ESTADOS.index(row["ESTADO"]) if row["ESTADO"] in ESTADOS else 0,key=f"est_{wk}")
            with e2: new_av=st.slider("Avance %",0,100,av_v,key=f"av_{wk}")
            with e3: new_resp=st.text_input("Responsable",value=row["RESPONSABLE"],key=f"rsp_{wk}")
            new_obs=st.text_area("Observación",value=row.get("OBSERVACIÓN",""),height=65,key=f"obs_{wk}")
            try:    fl_val=datetime.strptime(row.get("FECHA LÍMITE",""),"%d/%m/%Y").date() if row.get("FECHA LÍMITE","").strip() else None
            except: fl_val=None
            new_fl=st.date_input("Fecha límite",value=fl_val,key=f"fl_{wk}")
            b1,b2,b3,b4=st.columns(4)
            with b1:
                if st.button("💾 Guardar",key=f"sv_{wk}"):
                    df_new=st.session_state["df"].copy(); m=df_new["ID"]==row["ID"]
                    df_new.loc[m,"ESTADO"]=new_est; df_new.loc[m,"AVANCE"]=str(new_av)
                    df_new.loc[m,"RESPONSABLE"]=new_resp; df_new.loc[m,"OBSERVACIÓN"]=new_obs
                    if new_fl: df_new.loc[m,"FECHA LÍMITE"]=new_fl.strftime("%d/%m/%Y")
                    commit(df_new,f"{row['ID']} actualizada: {new_est}, {new_av}%"); st.rerun()
            with b2:
                if row["ESTADO"]!="Cerrada":
                    if st.button("✅ Cerrar",key=f"cl_{wk}"):
                        df_new=st.session_state["df"].copy(); m=df_new["ID"]==row["ID"]
                        df_new.loc[m,"ESTADO"]="Cerrada"; df_new.loc[m,"AVANCE"]="100"
                        df_new.loc[m,"FECHA CIERRE"]=datetime.now().strftime("%d/%m/%Y")
                        commit(df_new,f"{row['ID']} CERRADA"); st.rerun()
                else:
                    if st.button("🔄 Reabrir",key=f"ro_{wk}"):
                        df_new=st.session_state["df"].copy(); m=df_new["ID"]==row["ID"]
                        df_new.loc[m,"ESTADO"]="En curso"; df_new.loc[m,"FECHA CIERRE"]=""
                        commit(df_new,f"{row['ID']} reabierta"); st.rerun()
            with b3:
                if st.button("🟠 Bloquear",key=f"bl_{wk}"):
                    df_new=st.session_state["df"].copy(); df_new.loc[df_new["ID"]==row["ID"],"ESTADO"]="Bloqueada"
                    commit(df_new,f"{row['ID']} BLOQUEADA"); st.rerun()
            with b4:
                if st.button("🗑️ Eliminar",key=f"dl_{wk}"):
                    df_new=st.session_state["df"].copy(); df_new=df_new[df_new["ID"]!=row["ID"]].reset_index(drop=True)
                    commit(df_new,f"{row['ID']} ELIMINADA"); st.rerun()

# ── KANBAN ───────────────────────────────────────────────────────────────────
elif pagina=="🗂️ Kanban":
    st.markdown("## 🗂️ Tablero Kanban")
    fa_k=st.selectbox("Filtrar área",["Todas"]+sorted(df["ÁREA"].unique().tolist()),key="ka")
    fk=df if fa_k=="Todas" else df[df["ÁREA"]==fa_k]
    for col,(est,icon) in zip(st.columns(4),{"Sin iniciar":"🔴","En curso":"🟡","Bloqueada":"🟠","Cerrada":"🟢"}.items()):
        sub_k=fk[fk["ESTADO"]==est]; col.markdown(f"**{icon} {est}** `{len(sub_k)}`")
        col.markdown('<hr style="margin:5px 0;border-color:#eee">',unsafe_allow_html=True)
        for _,r in sub_k.iterrows():
            try:    avk=int(float(r.get("AVANCE","0") or 0))
            except: avk=0
            ac=AREA_COLOR.get(r["ÁREA"],"#888")
            col.markdown(f'<div style="background:white;border:0.5px solid #e0e0e0;border-radius:8px;padding:10px;margin-bottom:8px;border-top:3px solid {ac}">'
                f'<div style="font-size:.68rem;color:#999">{r["ID"]} · {r["ÁREA"]}</div>'
                f'<div style="font-size:.8rem;font-weight:600;color:#1B3A5C;margin:3px 0 4px;line-height:1.3">{r["TAREA"][:60]}{"…" if len(r["TAREA"])>60 else ""}</div>'
                f'<div style="font-size:.68rem;color:#aaa">{r["RESPONSABLE"]}</div>'
                f'<div style="height:3px;background:#eee;border-radius:2px;margin-top:5px"><div style="height:3px;background:{ac};width:{avk}%;border-radius:2px"></div></div>'
                f'<div style="font-size:.64rem;color:#bbb;text-align:right;margin-top:2px">{avk}%</div></div>',unsafe_allow_html=True)

# ── PROYECTOS ────────────────────────────────────────────────────────────────
elif pagina=="📊 Proyectos ZNI":
    st.markdown("## 📊 Proyectos ZNI — 35 Municipios")
    pdata=[("Dispac","Cartagena del Chaira",1118,"CRÍTICA"),("Dispac","Puerto Leguizamo",1154,"CRÍTICA"),
        ("Dispac","Miraflores",1057,"CRÍTICA"),("Sunco Energy","Linea Colectora",1122,"CRÍTICA"),
        ("Dispac","Tierralta",976,"ALTA"),("Dispac","Tolima",858,"ALTA"),("Dispac","Guainia",798,"ALTA"),
        ("IPSE","Michikai Norte y Centro",749,"ALTA"),("IPSE","Michikai Sur y Noreste",660,"ALTA"),
        ("Dispac","Unguia",651,"ALTA"),("Gensa","Solano",503,"ALTA"),("Gensa","Amazonas 659",491,"ALTA"),
        ("Dispac","Cumaribo",485,"ALTA"),("Gensa","Puerto Carreño",355,"MEDIA"),("Dispac","Morales",279,"MEDIA"),
        ("Gensa","Fundacion",274,"MEDIA"),("Dispac","Milan",254,"MEDIA"),("Dispac","Barrancas",235,"MEDIA"),
        ("Dispac","Riohacha",220,"MEDIA"),("Gensa","Amazonas 653",157,"MEDIA"),
        ("Cedenar","Puerto Asis",154,"MEDIA"),("Cedenar","Mocoa",151,"MEDIA"),
        ("CENS","Tibú",130,"MEDIA"),("Dispac","Paz de Ariporo",131,"MEDIA"),("Dispac","Maicao",153,"MEDIA"),
        ("Dispac","Zambrano",104,"BAJA"),("Cedenar","Orito",105,"BAJA"),("Cedenar","Valle del Guamuez",78,"BAJA"),
        ("ISA","El Copey",58,"BAJA"),("CENS","El Carmen",58,"BAJA"),("CENS","Sardinata",33,"BAJA"),
        ("CENS","El Tarra",35,"BAJA"),("CENS","Abrego",19,"BAJA"),("CENS","Teorama",3,"BAJA"),("CENS","Convención",4,"BAJA")]
    dp=pd.DataFrame(pdata,columns=["Proyecto","Municipio","Usuarios","Prioridad"])
    m1,m2,m3,m4=st.columns(4)
    m1.metric("Proyectos","7"); m2.metric("Municipios","35"); m3.metric("Usuarios totales",f"{dp['Usuarios'].sum():,}"); m4.metric("Críticos",str(len(dp[dp["Prioridad"]=="CRÍTICA"])))
    st.markdown("---"); cl2,cr2=st.columns(2)
    with cl2:
        fig_p=px.bar(dp.sort_values("Usuarios",ascending=True).tail(15),x="Usuarios",y="Municipio",orientation="h",color="Prioridad",
            color_discrete_map={"CRÍTICA":"#C0392B","ALTA":"#E67E22","MEDIA":"#27AE60","BAJA":"#2E6DA4"},title="Top 15 por usuarios")
        fig_p.update_layout(height=380,margin=dict(l=0,r=0,t=40,b=0),plot_bgcolor="white",paper_bgcolor="white",font=dict(family="Arial",size=11))
        st.plotly_chart(fig_p,use_container_width=True)
    with cr2:
        res=dp.groupby("Proyecto").agg(Municipios=("Municipio","count"),Usuarios=("Usuarios","sum")).reset_index().sort_values("Usuarios",ascending=False)
        st.dataframe(res,use_container_width=True,hide_index=True)

# ── NUEVA TAREA ──────────────────────────────────────────────────────────────
elif pagina=="➕ Nueva Tarea":
    st.markdown("## ➕ Nueva Tarea")
    with st.form("form_nueva",clear_on_submit=True):
        r1,r2=st.columns(2)
        with r1: nt_t=st.text_area("Título *",height=80); nt_a=st.selectbox("Área *",list(AREA_COLOR.keys())); nt_p=st.text_input("Proyecto/Municipio",placeholder="Todos / Cartagena del Chaira...")
        with r2: nt_r=st.text_input("Responsable *"); nt_pr=st.selectbox("Prioridad *",PRIORIDADES); nt_c=st.text_input("Categoría"); nt_fl=st.date_input("Fecha límite",value=None); nt_av=st.slider("Avance inicial %",0,100,0)
        nt_obs=st.text_area("Observaciones",height=70)
        if st.form_submit_button("✅ Crear tarea",type="primary"):
            if not nt_t.strip(): st.error("El título es obligatorio.")
            elif not nt_r.strip(): st.error("El responsable es obligatorio.")
            else:
                df_new=st.session_state["df"].copy(); nid=next_id(df_new); nr={c:"" for c in COLS}
                nr.update({"ID":nid,"ÁREA":nt_a,"PROYECTO":nt_p or "Todos","TAREA":nt_t.strip(),"CATEGORÍA":nt_c or "General",
                    "RESPONSABLE":nt_r.strip(),"FECHA INICIO":datetime.now().strftime("%d/%m/%Y"),
                    "FECHA LÍMITE":nt_fl.strftime("%d/%m/%Y") if nt_fl else "","PRIORIDAD":nt_pr,"AVANCE":str(nt_av),"ESTADO":"Sin iniciar","OBSERVACIÓN":nt_obs.strip()})
                df_new=pd.concat([df_new,pd.DataFrame([nr])],ignore_index=True)
                commit(df_new,f"{nid} creada: {nt_t[:45]} [{nt_a}]"); st.success(f"✅ Tarea {nid} creada."); st.rerun()

# ── HISTORIAL ────────────────────────────────────────────────────────────────
elif pagina=="📜 Historial":
    st.markdown("## 📜 Historial")
    hist=st.session_state.get("hist",[])
    if not hist: st.info("Sin actividad.")
    else:
        for h in hist: st.markdown(f'<div style="padding:7px 0;border-bottom:0.5px solid #eee;font-size:.86rem"><span style="color:#aaa;margin-right:10px">{h["ts"]}</span>{h["msg"]}</div>',unsafe_allow_html=True)
    st.markdown("---"); ec1,ec2=st.columns(2)
    with ec1: st.download_button("⬇️ CSV",df.to_csv(index=False).encode("utf-8-sig"),f"DISPOWER_{datetime.now().strftime('%Y%m%d')}.csv","text/csv")
    with ec2:
        b2=io.BytesIO(); df.to_excel(b2,index=False,engine="openpyxl")
        st.download_button("⬇️ Excel",b2.getvalue(),f"DISPOWER_{datetime.now().strftime('%Y%m%d')}.xlsx","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

# ── CONFIGURACIÓN ────────────────────────────────────────────────────────────
elif pagina=="⚙️ Configuración":
    st.markdown("## ⚙️ Configuración")
    cfg_now=get_cfg(); tok_c,_=get_access_token()
    if tok_c: st.markdown('<div class="bok">✅ Token obtenido. OneDrive conectado.</div>',unsafe_allow_html=True)
    else:      st.markdown('<div class="berr">❌ No se pudo obtener el token.</div>',unsafe_allow_html=True)
    if tok_c:
        ep=file_endpoint(); info=g_json(ep.replace("/content",""),tok_c)
        if info and info.get("name"): st.markdown(f'<div class="bok">✅ Archivo: <b>{info["name"]}</b> — {int(info.get("size",0)/1024)} KB</div>',unsafe_allow_html=True)
        else: st.markdown(f'<div class="bwarn">⚠️ Archivo no encontrado en: <code>{cfg_now.get("file_path","")}</code></div>',unsafe_allow_html=True)
    st.markdown("---"); st.markdown("### secrets.toml")
    st.markdown('<div class="binfo">Usa la sección <b>[microsoft]</b> — mismo formato que el otro app de la empresa.</div>',unsafe_allow_html=True)
    st.code("""[microsoft]
tenant_id     = "7df7bf9b-7611-4597-b13a-07ed6df6fac2"
client_id     = "2835aaaa-9a71-41df-bc88-d8d82d695bd3"
client_secret = "vOq8Q~UvGoC8vMZWIngA6LD1~LADMr8cKtuLbcu7"
file_path     = "/DISPOWER_Tareas_Streamlit.xlsx"
refresh_token = "1.AXEBm7_3..."   ← el mismo del otro app
""",language="toml")
    st.markdown("**`file_path`**: ruta del Excel en OneDrive desde la raíz. Ejemplos:")
    st.code('file_path = "/DISPOWER_Tareas_Streamlit.xlsx"          # en la raíz\nfile_path = "/Documentos/DISPOWER_Tareas_Streamlit.xlsx"\nfile_path = "/1. Escalamiento/DISPOWER_Tareas_Streamlit.xlsx"',language="toml")
    if st.button("🔍 Verificar conexión y archivo"):
        tok_t,_=get_access_token()
        if tok_t:
            st.success("✅ Token OK")
            ep2=file_endpoint(); i2=g_json(ep2.replace("/content",""),tok_t)
            if i2 and i2.get("name"): st.success(f"✅ Archivo: **{i2['name']}** — {int(i2.get('size',0)/1024)} KB")
            else:
                st.error(f"❌ Archivo no encontrado. Ruta: `{cfg_now.get('file_path','')}`")
                root=g_json("me/drive/root/children",tok_t)
                if root and root.get("value"):
                    st.markdown("**Archivos en la raíz de tu OneDrive:**")
                    st.write([i.get("name","") for i in root["value"]])
        else: st.error("❌ Error de autenticación.")
    st.markdown("---")
    if st.button("🗑️ Resetear sesión"):
        st.session_state.clear(); get_access_token.clear(); download_excel.clear(); st.rerun()
